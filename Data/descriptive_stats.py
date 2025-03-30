import pandas as pd
from docx import Document

# Load your CSV file
file_path = "/Users/bobi/Desktop/FIN 427/ML_In_Investing/Data/Final data 20250312_2300.csv"
columns = [
    'lag1mcreal', 'g01dyadj', 'g02esg', 'g03nibadj', 'g04fcfyadj',
    'g05rdsadj', 'g06_invpegadj', 'g07epadj', 'g08sadadj', 'g09shoadj',
    'g10shiadj', 'g11ret5adj', 'g12empadj', 'g13sueadj', 'g14erevadj'
]

# Read the data and compute descriptive statistics
df = pd.read_csv(file_path)
stats = df[columns].describe()

# Create a new Word document
doc = Document()

# Add a title for your table
doc.add_heading('Descriptive Statistics', 0)

# Add a table to the document
table = doc.add_table(rows=1, cols=len(stats.columns)+1)  # +1 for row names

# Add the header row
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'Statistic'
for i, column in enumerate(stats.columns, 1):
    hdr_cells[i].text = column

# Add the data rows
for stat_name, stat_values in stats.iterrows():
    row_cells = table.add_row().cells
    row_cells[0].text = stat_name
    for i, value in enumerate(stat_values, 1):
        row_cells[i].text = str(value)

# Save the Word document
output_path = '/Users/bobi/Desktop/FIN 427/ML_In_Investing/Data/Descriptive_Data.docx'
doc.save(output_path)

print(f"Table saved to {output_path}")




